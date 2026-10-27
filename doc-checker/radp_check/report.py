"""Excel report. The "Assessor agrees?" and "Assessor comment" columns are filled in by a person;
the `feedback` command then pulls them into the knowledge base."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .checker import Context, Result
from .knowledge import Knowledge

FILLS = {s: PatternFill("solid", fgColor=c) for s, c in
         {"OK": "C6EFCE", "INFO": "DDEBF7", "WARN": "FFEB9C", "FAIL": "FFC7CE"}.items()}
STATUS_LABEL = {"OK": "OK", "INFO": "INFO", "WARN": "WARNING", "FAIL": "FAIL"}
ORDER = {"FAIL": 0, "WARN": 1, "INFO": 2, "OK": 3}
CHECK_COLUMNS = ["LoR row", "Serial No. (LoR)", "Description (LoR)", "Category", "Check", "LoR inspection date",
                 "LoR expiry date", "Result", "Reason", "File", "Page", "Document date", "Document valid until",
                 "Issued by", "Rule", "Rule status", "Requirement source", "Assessor agrees?", "Assessor comment"]
FEEDBACK_COL = CHECK_COLUMNS.index("Assessor agrees?") + 1


def _header(ws, columns, widths=None):
    ws.append(columns)
    for i, cell in enumerate(ws[1], start=1):
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = (widths or {}).get(i, 14)
    ws.freeze_panes = "A2"


def _d(d) -> str:
    return d.strftime("%d-%b-%Y") if d else ""


def write_report(path: Path, kn: Knowledge, ctx: Context, result: Result, provider: str) -> Path:
    wb = Workbook()
    counts = result.counts()
    info = result.info
    in_scope = sum(s.in_scope for s in result.scope)

    ws = wb.active
    ws.title = "Summary"
    overall = ("RETURN TO CONTRACTOR" if counts.get("FAIL") else
               "READY FOR REVIEW (with remarks)" if counts.get("WARN") or result.unrecognized else "READY FOR REVIEW")
    start, end = ctx.period
    rows = [
        ("Request", ctx.request_id), ("Service provider", info.provider or "-"),
        ("Field / well", info.field or "-"), ("Type of job", info.job_type or "-"),
        ("LoR date of issue", _d(info.issue_date) or "-"),
        ("Discipline checked", ctx.discipline or "all"),
        ("Validity checked for", f"{_d(start)} to {_d(end)}" if ctx.work_start else f"{_d(start)} (no work dates given)"),
        ("Conditions", ", ".join(f"{k}={v}" for k, v in ctx.flags.items()) or "-"),
        ("Checked", datetime.now().strftime("%d-%b-%Y %H:%M")), ("Document reader", provider),
        ("", ""),
        ("LoR equipment lines", f"{len(result.scope)} ({in_scope} in scope)"),
        ("Checks evaluated", len(result.findings)),
        ("OK", counts.get("OK", 0)), ("Warning", counts.get("WARN", 0)), ("Fail", counts.get("FAIL", 0)),
        ("Info", counts.get("INFO", 0)), ("Unrecognised items", len(result.unrecognized)),
        ("Assignment remarks", sum(bool(s.note) and "Assigned" in s.note for s in result.scope)),
        ("", ""), ("Preliminary status", overall),
        ("", ""), ("IMPORTANT", "This is a pre-screening. The assessor makes the decision."),
    ]
    if any(f.rule and f.rule.get("status") != "approved" for f in result.findings):
        rows.append(("", "UNAPPROVED (draft) rules were used - result is for testing only."))
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 70
    for cell in ws["A"]:
        cell.font = Font(bold=True)

    ws = wb.create_sheet("Check")
    _header(ws, CHECK_COLUMNS, {3: 28, 9: 60, 10: 30, 14: 24, 17: 24, 19: 30})
    for f in sorted(result.findings, key=lambda f: (ORDER[f.status], f.serial)):
        c, rule, row = f.cert or {}, f.rule or {}, f.lor_row
        src = rule.get("source") or {}
        ws.append([
            row.row if row else "", f.serial, f.description,
            kn.equipment.get(f.category, {}).get("name", f.category or ""),
            kn.doc_name(f.check_type) if f.check_type else (row.check_raw if row else ""),
            _d(row.inspection_date) if row else "",
            (_d(row.expiry_date) or row.expiry_text) if row else "",
            STATUS_LABEL[f.status], "; ".join(f.reasons),
            Path(c["file"]).name if c.get("file") else "", c.get("page") or "",
            c.get("issue_date") or "", f.valid_until.isoformat() if f.valid_until else (c.get("expiry_date") or ""),
            c.get("issuer") or "", rule.get("id", ""), rule.get("status", ""),
            " ".join(x for x in (src.get("document"), src.get("clause"), src.get("revision")) if x), "", "",
        ])
        r = ws.max_row
        ws.cell(r, CHECK_COLUMNS.index("Result") + 1).fill = FILLS[f.status]
        for col in (3, 9, 10):
            ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    letter = ws.cell(1, FEEDBACK_COL).column_letter
    dv.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Scope")
    _header(ws, ["LoR row", "Serial No.", "Description", "Check", "Responsible unit (LoR)", "Category",
                 "Discipline", "In scope", "Remark"], {3: 30, 9: 60})
    for s in sorted(result.scope, key=lambda s: (not s.note, s.row.row)):
        ws.append([s.row.row, s.row.serial, s.row.description, s.row.check_raw, s.row.unit,
                   kn.equipment.get(s.row.category, {}).get("name", ""), s.discipline or "?",
                   "yes" if s.in_scope else "no", s.note])
        if s.note:
            ws.cell(ws.max_row, 9).fill = FILLS["INFO"]
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Unrecognised")
    _header(ws, ["Type", "Item", "Reason"], {1: 20, 2: 45, 3: 80})
    for u in result.unrecognized:
        ws.append([u["what"], u["name"], u["reason"]])

    ws = wb.create_sheet("Extracted")
    cols = ["cert_id", "file", "page", "doc_type", "certificate_number", "serial_numbers",
            "equipment_description", "issue_date", "expiry_date", "issuer", "notes"]
    _header(ws, cols, {2: 40, 6: 40, 7: 30, 10: 28, 11: 40})
    for c in result.certs:
        ws.append([Path(c[k]).name if k == "file" else
                   ", ".join(c[k]) if k == "serial_numbers" else c.get(k) for k in cols])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
