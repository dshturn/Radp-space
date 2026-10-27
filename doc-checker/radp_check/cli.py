"""Командная строка: python -m radp_check <команда> ... (подробно — README.md)."""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook

from . import samples
from .checker import Context, run_check
from .extract import extract_all, get_extractor
from .ingest import WORK_DIR_NAME, collect
from .knowledge import RULE_STATUSES, Knowledge
from .lor import read_lor
from .report import CHECK_COLUMNS, write_report
from .util import parse_date


def _flag(value: str) -> tuple[str, object]:
    key, _, raw = value.partition("=")
    return key.strip(), yaml.safe_load(raw) if raw else True


# ---------------- check ----------------

def cmd_check(args, kn: Knowledge) -> int:
    pkg = collect(Path(args.package), args.lor)
    req = yaml.safe_load(pkg.request_file.read_text(encoding="utf-8")) if pkg.request_file else {}
    flags = dict(req.get("flags") or {})
    flags.update(dict(_flag(f) for f in args.flag or []))
    ctx = Context(
        request_id=args.request_id or req.get("request_id") or Path(args.package).name,
        service_line=args.service_line or req.get("service_line") or "",
        work_start=parse_date(args.work_start or req.get("work_start")),
        work_end=parse_date(args.work_end or req.get("work_end") or args.work_start or req.get("work_start")),
        flags=flags,
    )
    missing = [n for n, v in (("--service-line", ctx.service_line), ("--work-start", ctx.work_start)) if not v]
    if missing:
        print(f"Не хватает контекста заявки: {', '.join(missing)} (или request.yaml в папке пакета)")
        return 2
    if not pkg.lor:
        print("В пакете не найден LoR (.xlsx). Укажите явно: --lor <файл>")
        return 2

    print(f"Заявка {ctx.request_id} | {ctx.service_line} | работы {ctx.work_start} – {ctx.work_end} | {flags or ''}")
    items, lor_problems = read_lor(pkg.lor, kn)
    print(f"LoR: {pkg.lor.name} — {len(items)} позиций")
    for p in lor_problems:
        print(f"  ! {p}")

    print(f"Документы: {len(pkg.documents)} файлов, чтение: {args.provider}")
    extractor = get_extractor(args.provider)
    certs, failures = extract_all(pkg.documents, extractor, kn, Path(args.package) / WORK_DIR_NAME / "cache")

    result = run_check(kn, ctx, items, certs, pkg.skipped + failures, args.approved_only)
    out = Path(args.out) if args.out else (Path(args.package) / WORK_DIR_NAME /
                                           f"report_{ctx.request_id}_{datetime.now():%Y%m%d_%H%M}.xlsx")
    write_report(out, kn, ctx, result, args.provider)

    c = result.counts()
    print(f"\nИтог: OK {c.get('OK', 0)} | ВНИМАНИЕ {c.get('WARN', 0)} | НЕ СООТВ. {c.get('FAIL', 0)} | "
          f"нераспознанное {len(result.unrecognized)}")
    print(f"Отчёт: {out}")
    return 0


# ---------------- обучение: правила ----------------

def cmd_rules_list(args, kn: Knowledge) -> int:
    for r in kn.rules:
        if args.service_line and r.get("service_line") != args.service_line:
            continue
        when = ", ".join(f"{k}={v}" for k, v in (r.get("when") or {}).items())
        print(f"{r['id']:<12} [{r.get('status', 'draft'):<8}] {r['equipment']:<18} -> {r['document']:<26} "
              f"max {r.get('max_age_months') or '-'} мес."
              f"{'  3rd party' if r.get('third_party_required') else ''}{'  если ' + when if when else ''}")
    return 0


def cmd_rules_add(args, kn: Knowledge) -> int:
    if kn.rule(args.id):
        print(f"Правило {args.id} уже есть. Чтобы изменить — отключите его и добавьте новое с другим id.")
        return 1
    rule = {
        "id": args.id, "service_line": args.service_line, "equipment": args.equipment,
        "document": args.document, "when": dict(_flag(f) for f in args.when or []),
        "max_age_months": args.max_age, "serial_must_match": not args.no_serial_match,
        "third_party_required": args.third_party, "status": "draft",
        "source": {"document": args.source_doc or "PLACEHOLDER", "clause": args.clause or "",
                   "revision": args.revision or ""},
        "note": args.note or "",
    }
    kn.rules.append(rule)
    problems = [p for p in kn.validate() if p.startswith(args.id)]
    if problems:
        print("Правило не добавлено:\n  " + "\n  ".join(problems))
        return 1
    kn.save_rules()
    kn.log("rule_add", rule=rule)
    print(f"Добавлено {args.id} (status: draft). Утвердить: rules status {args.id} approved")
    return 0


def cmd_rules_status(args, kn: Knowledge) -> int:
    rule = kn.rule(args.id)
    if not rule:
        print(f"Нет правила {args.id}")
        return 1
    if args.status == "approved" and (rule.get("source") or {}).get("document") in (None, "", "PLACEHOLDER"):
        print("Нельзя утвердить правило без источника требования (документ/пункт/ревизия).")
        return 1
    old = rule.get("status")
    rule["status"] = args.status
    kn.save_rules()
    kn.log("rule_status", id=args.id, old=old, new=args.status, reason=args.reason)
    print(f"{args.id}: {old} -> {args.status}")
    return 0


# ---------------- обучение: словари ----------------

def _table(kn: Knowledge, kind: str):
    return (kn.equipment, kn.save_equipment) if kind == "equipment" else (kn.doc_types, kn.save_doc_types)


def cmd_alias_add(args, kn: Knowledge) -> int:
    table, save = _table(kn, args.kind)
    if args.key not in table:
        print(f"Нет ключа '{args.key}'. Существующие: {', '.join(table)}")
        return 1
    aliases = table[args.key].setdefault("aliases", [])
    if args.alias.lower() in (a.lower() for a in aliases):
        print("Такой синоним уже есть")
        return 0
    aliases.append(args.alias)
    save()
    kn.log("alias_add", kind=args.kind, key=args.key, alias=args.alias)
    print(f"{args.kind}/{args.key}: добавлен синоним '{args.alias}'")
    return 0


def cmd_category_add(args, kn: Knowledge) -> int:
    table, save = _table(kn, args.kind)
    if args.key in table:
        print(f"'{args.key}' уже существует")
        return 1
    table[args.key] = {"name": args.name, "aliases": args.alias or []}
    save()
    kn.log("category_add", kind=args.kind, key=args.key, name=args.name, aliases=args.alias or [])
    print(f"Добавлено {args.kind}/{args.key}")
    return 0


def cmd_issuer_add(args, kn: Knowledge) -> int:
    if kn.issuer_approved(args.name):
        print("Уже в списке")
        return 0
    kn.issuers.append(args.name)
    kn.save_issuers()
    kn.log("issuer_add", name=args.name, reason=args.reason)
    print(f"Добавлено: {args.name}")
    return 0


# ---------------- обратная связь, экспорт, проверка ----------------

def cmd_feedback(args, kn: Knowledge) -> int:
    ws = load_workbook(args.report, read_only=True)["Проверка"]
    rows = list(ws.iter_rows(values_only=True))
    idx = {name: i for i, name in enumerate(rows[0])}
    agree, disagree = 0, []
    for row in rows[1:]:
        verdict = str(row[idx["Ассессор: согласен?"]] or "").strip().lower()
        if verdict == "да":
            agree += 1
        elif verdict == "нет":
            disagree.append({k: row[idx[k]] for k in CHECK_COLUMNS if k in idx})
    with open(kn.dir / "feedback.jsonl", "a", encoding="utf-8") as f:
        for d in disagree:
            f.write(json.dumps({"ts": datetime.now().isoformat(timespec="seconds"),
                                "report": Path(args.report).name, **d}, ensure_ascii=False, default=str) + "\n")
    total = agree + len(disagree)
    if not total:
        print("В отчёте не заполнена колонка «Ассессор: согласен?»")
        return 1
    print(f"Размечено {total}: согласен {agree}, не согласен {len(disagree)} "
          f"(точность {agree / total:.0%}). Несогласия сохранены в knowledge/feedback.jsonl")
    for rule_id, n in Counter(d["Правило"] or "(без правила)" for d in disagree).most_common():
        print(f"  {rule_id}: {n} — пересмотрите правило или синонимы")
        for d in disagree:
            if (d["Правило"] or "(без правила)") == rule_id and d.get("Комментарий ассессора"):
                print(f"      «{d['Комментарий ассессора']}»")
    return 0


def cmd_export(args, kn: Knowledge) -> int:
    """Экспорт базы знаний в Excel — формат для импорта в SharePoint List."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"
    ws.append(["RuleID", "ServiceLine", "Equipment", "Document", "Conditions", "MaxAgeMonths",
               "SerialMustMatch", "ThirdPartyRequired", "Status", "SourceDocument", "Clause", "Revision", "Note"])
    for r in kn.rules:
        s = r.get("source") or {}
        ws.append([r["id"], r["service_line"], r["equipment"], r["document"],
                   json.dumps(r.get("when") or {}), r.get("max_age_months"), r.get("serial_must_match"),
                   r.get("third_party_required"), r.get("status"), s.get("document"), s.get("clause"),
                   s.get("revision"), r.get("note")])
    for title, table in (("EquipmentCategories", kn.equipment), ("DocumentTypes", kn.doc_types)):
        ws = wb.create_sheet(title)
        ws.append(["Key", "Name", "Aliases"])
        for k, v in table.items():
            ws.append([k, v.get("name"), "; ".join(v.get("aliases", []))])
    ws = wb.create_sheet("ApprovedIssuers")
    ws.append(["Name"])
    for i in kn.issuers:
        ws.append([i])
    wb.save(args.out)
    print(f"Экспортировано в {args.out}")
    return 0


def cmd_validate(args, kn: Knowledge) -> int:
    problems = kn.validate()
    print("\n".join(problems) if problems else
          f"OK: {len(kn.rules)} правил, {len(kn.equipment)} категорий, {len(kn.doc_types)} типов документов")
    return 1 if problems else 0


def cmd_generate(args, kn: Knowledge) -> int:
    pkg = samples.generate(Path(args.dir))
    print(f"Фейковый пакет: {pkg}\nЭталон дефектов: {Path(args.dir) / 'expected.json'}")
    print(f"Проверить: python -m radp_check check \"{pkg}\"")
    return 0


# ---------------- argparse ----------------

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="radp_check", description="Предварительная проверка пакетов readiness")
    p.add_argument("--knowledge", help="Папка базы знаний (по умолчанию ./knowledge)")
    sub = p.add_subparsers(dest="cmd", required=True)

    c = sub.add_parser("check", help="Проверить пакет")
    c.add_argument("package")
    c.add_argument("--provider", choices=["mock", "claude", "openai"], default="mock")
    c.add_argument("--request-id")
    c.add_argument("--service-line")
    c.add_argument("--work-start")
    c.add_argument("--work-end")
    c.add_argument("--flag", action="append", help="Условие работ, например h2s_service=true")
    c.add_argument("--lor", help="Путь к LoR, если не определился автоматически")
    c.add_argument("--approved-only", action="store_true", help="Только утверждённые правила")
    c.add_argument("--out", help="Куда сохранить отчёт")
    c.set_defaults(fn=cmd_check)

    g = sub.add_parser("generate-sample", help="Создать фейковый пакет с дефектами")
    g.add_argument("dir")
    g.set_defaults(fn=cmd_generate)

    r = sub.add_parser("rules", help="Свод правил").add_subparsers(dest="sub", required=True)
    rl = r.add_parser("list")
    rl.add_argument("--service-line")
    rl.set_defaults(fn=cmd_rules_list)
    ra = r.add_parser("add")
    ra.add_argument("--id", required=True)
    ra.add_argument("--service-line", required=True)
    ra.add_argument("--equipment", required=True, help="ключ категории из equipment.yaml")
    ra.add_argument("--document", required=True, help="ключ типа из document_types.yaml")
    ra.add_argument("--max-age", type=int, help="Максимальный возраст документа, мес.")
    ra.add_argument("--third-party", action="store_true", help="Требуется одобренная сторонняя организация")
    ra.add_argument("--no-serial-match", action="store_true", help="Не требовать совпадения серийника")
    ra.add_argument("--when", action="append", help="Условие, например h2s_service=true")
    ra.add_argument("--source-doc")
    ra.add_argument("--clause")
    ra.add_argument("--revision")
    ra.add_argument("--note")
    ra.set_defaults(fn=cmd_rules_add)
    rs = r.add_parser("status")
    rs.add_argument("id")
    rs.add_argument("status", choices=RULE_STATUSES)
    rs.add_argument("--reason")
    rs.set_defaults(fn=cmd_rules_status)

    a = sub.add_parser("alias", help="Синонимы").add_subparsers(dest="sub", required=True)
    aa = a.add_parser("add")
    aa.add_argument("kind", choices=["equipment", "document"])
    aa.add_argument("key")
    aa.add_argument("alias")
    aa.set_defaults(fn=cmd_alias_add)

    k = sub.add_parser("category", help="Новые категории оборудования / типы документов").add_subparsers(
        dest="sub", required=True)
    ka = k.add_parser("add")
    ka.add_argument("kind", choices=["equipment", "document"])
    ka.add_argument("key")
    ka.add_argument("name")
    ka.add_argument("--alias", action="append")
    ka.set_defaults(fn=cmd_category_add)

    i = sub.add_parser("issuer", help="Одобренные организации").add_subparsers(dest="sub", required=True)
    ia = i.add_parser("add")
    ia.add_argument("name")
    ia.add_argument("--reason")
    ia.set_defaults(fn=cmd_issuer_add)

    f = sub.add_parser("feedback", help="Забрать разметку ассессора из отчёта")
    f.add_argument("report")
    f.set_defaults(fn=cmd_feedback)

    e = sub.add_parser("export", help="Экспорт базы знаний в Excel (для SharePoint)")
    e.add_argument("out")
    e.set_defaults(fn=cmd_export)

    v = sub.add_parser("validate", help="Проверить целостность базы знаний")
    v.set_defaults(fn=cmd_validate)
    return p


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args = build_parser().parse_args(argv)
    kn = Knowledge.load(Path(args.knowledge) if args.knowledge else None)
    return args.fn(args, kn)
