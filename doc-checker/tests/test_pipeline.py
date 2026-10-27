import json
import shutil
from datetime import date

from openpyxl import load_workbook

from radp_check.cli import main
from radp_check.knowledge import DEFAULT_DIR, Knowledge
from radp_check.util import near_miss, parse_lor_date

STATUS = {"OK": "OK", "WARN": "WARNING", "FAIL": "FAIL"}
DOC_NAMES = {k: v["name"] for k, v in Knowledge.load().doc_types.items()}


def _setup(tmp_path):
    kn_dir = tmp_path / "knowledge"
    shutil.copytree(DEFAULT_DIR, kn_dir)
    main(["--knowledge", str(kn_dir), "generate-sample", str(tmp_path / "demo")])
    return kn_dir, tmp_path / "demo" / "package"


def _check(kn_dir, pkg, out, *extra):
    assert main(["--knowledge", str(kn_dir), "check", str(pkg), "--out", str(out), *extra]) == 0
    wb = load_workbook(out)
    rows = list(wb["Check"].iter_rows(values_only=True))
    idx = {n: i for i, n in enumerate(rows[0])}
    findings = {f"{r[idx['Serial No. (LoR)']]}|{r[idx['Check']]}": (r[idx["Result"]], r[idx["Reason"]] or "")
                for r in rows[1:]}
    scope = list(wb["Scope"].iter_rows(values_only=True))
    unrecognized = [r[1] for r in list(wb["Unrecognised"].iter_rows(values_only=True))[1:]]
    return wb, findings, scope, unrecognized


def _key(k):
    serial, doc = k.split("|")
    return f"{serial}|{DOC_NAMES[doc]}"


def test_sample_package_defects_found(tmp_path):
    kn_dir, pkg = _setup(tmp_path)
    expected = json.loads((tmp_path / "demo" / "expected.json").read_text(encoding="utf-8"))
    _, findings, scope, unrecognized = _check(kn_dir, pkg, tmp_path / "r.xlsx")

    assert set(findings) == {_key(k) for k in expected["findings"]}
    for k, (status, reason_part) in expected["findings"].items():
        got_status, got_reason = findings[_key(k)]
        assert got_status == STATUS[status], (k, got_reason)
        assert reason_part in got_reason, (k, got_reason)
    for k, parts in expected["also_in_reason"].items():
        for part in parts:
            assert part in findings[_key(k)][1], (k, part)
    remarks = {r[1] for r in scope[1:] if r[8] and "Assigned" in r[8]}
    assert remarks == set(expected["assignment_remarks"])
    for name in expected["unrecognized_contains"]:
        assert any(name in u for u in unrecognized), name
    assert not any("MTC" in u or "Mill" in u for u in unrecognized)  # supporting documents are not noise


def test_lifting_out_of_scope_for_engineering(tmp_path):
    kn_dir, pkg = _setup(tmp_path)
    _, findings, scope, _ = _check(kn_dir, pkg, tmp_path / "r.xlsx")
    assert not any(k.split("|")[0] in ("SL-1483", "SH-6471", "SW-2289", "MGB-019") for k in findings)
    in_scope = {r[1]: r[7] for r in scope[1:]}
    assert in_scope["SL-1483"] == "no" and in_scope["WX-201"] == "yes"
    _, findings_all, _, _ = _check(kn_dir, pkg, tmp_path / "r_all.xlsx", "--discipline", "all")
    assert any(k.startswith("SL-1483|") for k in findings_all)


def test_learning_loop(tmp_path):
    kn_dir, pkg = _setup(tmp_path)
    run = lambda *a: main(["--knowledge", str(kn_dir), *a])

    # Approving the in-house test company clears the third-party warnings
    assert run("issuer", "add", "Fictitious Wireline Services") == 0
    _, findings, _, _ = _check(kn_dir, pkg, tmp_path / "r2.xlsx")
    assert findings[_key("WX-201|hydrostatic_test")][0] == "OK"

    # Mapping a unit code changes the discipline of unknown equipment
    assert run("alias", "add", "unit", "engineering", "HSE") == 0
    _, findings, _, _ = _check(kn_dir, pkg, tmp_path / "r3.xlsx")
    assert any(k.startswith("MGB-019|") for k in findings)

    # New category requires a discipline; a rule cannot be approved without a source
    assert run("category", "add", "equipment", "ptu", "Pressure test unit", "--alias", "ptu") == 1
    assert run("category", "add", "equipment", "ptu", "Pressure test unit", "--alias", "ptu",
               "--discipline", "engineering") == 0
    assert run("rules", "status", "PCE-BOP-001", "approved") == 1

    actions = [json.loads(l)["action"] for l in (kn_dir / "history.jsonl").read_text(encoding="utf-8").splitlines()]
    assert actions == ["issuer_add", "alias_add", "category_add"]
    assert Knowledge.load(kn_dir).validate() == []


def test_feedback_collected(tmp_path):
    kn_dir, pkg = _setup(tmp_path)
    wb, _, _, _ = _check(kn_dir, pkg, tmp_path / "r.xlsx")
    ws = wb["Check"]
    header = [c.value for c in ws[1]]
    col = header.index("Assessor agrees?") + 1
    comment = header.index("Assessor comment") + 1
    ws.cell(2, col, "Yes")
    ws.cell(3, col, "No")
    ws.cell(3, comment, "Contractor has sent a corrected certificate")
    wb.save(tmp_path / "r.xlsx")

    assert main(["--knowledge", str(kn_dir), "feedback", str(tmp_path / "r.xlsx")]) == 0
    lines = (kn_dir / "feedback.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(lines) == 1 and "corrected certificate" in lines[0]


def test_lor_dates_and_serials():
    assert parse_lor_date(46050) == (date(2026, 1, 28), None)
    assert parse_lor_date("16/3/2023")[0] == date(2023, 3, 16)
    assert parse_lor_date("-") == (None, None)
    assert parse_lor_date("until major repair")[0] is None
    d, remark = parse_lor_date("02/11/2026")
    assert d == date(2026, 2, 11) and "Ambiguous" in remark
    assert near_miss("EMX190037", "EMX19C037")
    assert not near_miss("WB094", "WB095X1")
    assert not near_miss("WB-1", "WB-2")  # too short to call a typo
